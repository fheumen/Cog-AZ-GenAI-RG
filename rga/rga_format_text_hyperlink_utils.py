import boto3
from io import BytesIO
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# -------- Minimal helpers --------
def highlight_run(run, color_name="green"):
    rPr = run._element.get_or_add_rPr()
    old = rPr.find(qn("w:highlight"))
    if old is not None:
        rPr.remove(old)
    hl = OxmlElement("w:highlight")
    hl.set(qn("w:val"), color_name)  # Word expects named colors like 'green'
    rPr.append(hl)

def clone_run_format(source_run, target_run):
    target_run.bold = source_run.bold
    target_run.italic = source_run.italic
    target_run.underline = source_run.underline
    if source_run.font.name:
        target_run.font.name = source_run.font.name
    if source_run.font.size:
        target_run.font.size = source_run.font.size
    if source_run.font.color and source_run.font.color.rgb:
        target_run.font.color.rgb = source_run.font.color.rgb

def should_replace(placeholder, value):
    return isinstance(placeholder, str) and isinstance(value, str) and value.strip()

# -------- Hyperlink-safe replace + highlight (namespace-agnostic XPath) --------
def highlight_subtext_in_hyperlink(hlink, substring, color_name="green"):
    """
    Highlight the first occurrence of substring within this hyperlink element without breaking the link.
    Uses .xpath() with local-name() predicates only.
    """
    # Iterate runs within hyperlink
    for r in hlink.xpath('.//*[local-name()="r"]'):
        # Get the first w:t under this run
        t_nodes = r.xpath('.//*[local-name()="t"]')
        if not t_nodes:
            continue
        t = t_nodes[0]
        if t is None or not t.text:
            continue
        text = t.text
        if substring not in text:
            continue

        before, match, after = text.partition(substring)

        if not before:
            # Match at the start: put match in current run and highlight it
            t.text = match
            # Ensure rPr exists (create with namespace-qualified tag)
            rPr_nodes = r.xpath('./*[local-name()="rPr"]')
            if rPr_nodes:
                rPr = rPr_nodes[0]
            else:
                rPr = OxmlElement("w:rPr")
                r.insert(0, rPr)
            hl = OxmlElement("w:highlight")
            hl.set(qn("w:val"), color_name)
            rPr.append(hl)
            if after:
                after_run = OxmlElement("w:r")
                after_t = OxmlElement("w:t")
                after_t.text = after
                after_run.append(after_t)
                r.addnext(after_run)
            break

        # There is a 'before' segment: keep it in current run
        t.text = before

        # Insert a new highlighted run for the match
        match_run = OxmlElement("w:r")
        match_rPr = OxmlElement("w:rPr")
        hl = OxmlElement("w:highlight")
        hl.set(qn("w:val"), color_name)
        match_rPr.append(hl)
        match_t = OxmlElement("w:t")
        match_t.text = match
        match_run.append(match_rPr)
        match_run.append(match_t)
        r.addnext(match_run)

        # Insert an 'after' run if any
        if after:
            after_run = OxmlElement("w:r")
            after_t = OxmlElement("w:t")
            after_t.text = after
            after_run.append(after_t)
            match_run.addnext(after_run)

        break  # only first occurrence per hyperlink

def handle_hyperlink_placeholders(paragraph, replacements, highlight_color="green"):
    # Find all hyperlink elements in the paragraph
    for hlink in paragraph._p.xpath('.//*[local-name()="hyperlink"]'):
        # Iterate text nodes within the hyperlink
        for t in hlink.xpath('.//*[local-name()="t"]'):
            if t is None or not t.text:
                continue
            txt = t.text
            for placeholder, value in replacements.items():
                if not should_replace(placeholder, value):
                    continue
                if placeholder in txt:
                    before, _, after = txt.partition(placeholder)
                    t.text = before + value + after
                    # Highlight only the inserted value within the hyperlink
                    highlight_subtext_in_hyperlink(hlink, value, color_name=highlight_color)
                    break  # next text node

# -------- General paragraph replacement (non-hyperlink text) --------
def replace_paragraph_text(paragraph, replacements, highlight_color="green"):
    if not paragraph.runs:
        return
    full_text = "".join(run.text for run in paragraph.runs)
    if not any(p in full_text for p in replacements):
        return

    # Map characters to runs (to preserve formatting when reconstructing)
    char_to_run = []
    for run in paragraph.runs:
        char_to_run.extend([run] * len(run.text))

    segments = []  # (text, source_run, is_highlighted)
    i = 0
    while i < len(full_text):
        matched = False
        for placeholder, value in sorted(replacements.items(), key=lambda x: -len(x[0])):
            if should_replace(placeholder, value) and full_text.startswith(placeholder, i):
                src_run = char_to_run[i] if i < len(char_to_run) else paragraph.runs[0]
                segments.append((value, src_run, True))
                i += len(placeholder)
                matched = True
                break
        if not matched:
            src_run = char_to_run[i] if i < len(char_to_run) else paragraph.runs[0]
            if segments and segments[-1][1] == src_run and not segments[-1][2]:
                segments[-1] = (segments[-1][0] + full_text[i], src_run, False)
            else:
                segments.append((full_text[i], src_run, False))
            i += 1

    # Clear and rebuild paragraph text
    for run in paragraph.runs:
        run.text = ""
    for text, src_run, is_hl in segments:
        nr = paragraph.add_run(text)
        clone_run_format(src_run, nr)
        if is_hl:
            highlight_run(nr, color_name=highlight_color)

# -------- Document-level orchestration --------
def apply_replacements_to_document(doc: Document, replacements: dict, highlight_color="green"):
    # Hyperlink-safe pass first
    for para in doc.paragraphs:
        handle_hyperlink_placeholders(para, replacements, highlight_color)
    # General text pass
    for para in doc.paragraphs:
        replace_paragraph_text(para, replacements, highlight_color)
    # Tables (do both passes per cell)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    handle_hyperlink_placeholders(para, replacements, highlight_color)
                for para in cell.paragraphs:
                    replace_paragraph_text(para, replacements, highlight_color)

# -------- S3 wrapper --------
def sitev_insert_text_in_report(bucket: str, template_key: str, output_key: str, replacements: dict, highlight_color="green"):
    s3 = boto3.client("s3")
    obj = s3.get_object(Bucket=bucket, Key=template_key)
    template_bytes = obj["Body"].read()

    doc = Document(BytesIO(template_bytes))
    apply_replacements_to_document(doc, replacements, highlight_color=highlight_color)

    out = BytesIO()
    doc.save(out)
    out.seek(0)
    s3.put_object(Bucket=bucket, Key=output_key, Body=out.read())

# # -------- Example --------
# if __name__ == "__main__":
#     bucket_name = "azcdi-us-ops-report-ds-dev"
#     input_key = "inputs/sv_docx_templates/Report 8 - Resin Lifetime Process Validation Template.docx"
#     output_key = "output_sv_2/sitev-ResinLifetimeAll-khld677_test2.docx"

#     data = {
#         "<Resin_Type>": "MabSelect™ SuRe™ (Cytiva) Protein A",
#         "<Molecule_Name>": "Tozorakimab",
#         "<Chromatography_Type>": "Protein A",
#         "<Building_Scale>": "15,000L",
#         "<Molecule_Number>": "MEDI3506",
#         "<Building_Number>": "B633",
#     }

#     generate_report_from_s3_template(
#         bucket=bucket_name,
#         template_key=input_key,
#         output_key=output_key,
#         replacements=data,
#         highlight_color="green",
#     )