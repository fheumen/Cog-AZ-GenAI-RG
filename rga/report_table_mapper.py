from io import BytesIO
import pandas as pd
from docx import Document
import re
import unicodedata
import logging
from typing import Optional
import io
import boto3
import docx
from docx import Document as Document_docx
from loguru import logger

# Constants
PARTIAL_MATCH_HEADERS = [
    "validation acceptance criteria", "Deviation"
]
VAC_HEADER_START = "validation acceptance criteria"
VAC_HEADER_END = "meets validation acceptance criteria (yes/no)"

# ----------------------------------------------------------------------
# Constants for Conditional Filtering
# ----------------------------------------------------------------------
PLACEHOLDER_FILTER_MAP = {
    "table 3": "vial thaw and inoculum expansion",
    "table 4": "seed bioreactors",
}
FILTER_COLUMN_NAME = "Unit Op Table"

s3 = boto3.client('s3')

# ------------------------------------------------------------------------------
# Utility Functions
# ------------------------------------------------------------------------------

def load_and_conditionally_filter_excel_sheet(bucket_name: str,
                                              xls_s3_key: str,
                                              sheet_name: str,
                                              placeholder: str) -> pd.DataFrame:
    """
    Loads an Excel sheet from S3 and applies conditional filtering based on the placeholder text.

    If the placeholder contains known keywords (like 'table 3' or 'table 4'), the corresponding filter
    value is used to filter the 'Unit Op Table' column in the Excel sheet.

    Args:
        bucket_name (str): S3 bucket name
        xls_s3_key (str): Excel file S3 key
        sheet_name (str): Sheet name to load
        placeholder (str): Placeholder text in DOCX

    Returns:
        pd.DataFrame: Loaded and possibly filtered DataFrame
    """
    # Load Excel sheet from S3
    response = s3.get_object(Bucket=bucket_name, Key=xls_s3_key)
    sheet_bytes = response['Body'].read()
    df = pd.read_excel(io.BytesIO(sheet_bytes), sheet_name=sheet_name)

    # Normalize placeholder for matching
    placeholder_norm = placeholder.lower()

    # Determine if filtering is required
    filter_value = None
    for key, val in PLACEHOLDER_FILTER_MAP.items():
        if key in placeholder_norm:
            filter_value = val
            break

    # Apply filtering if matched
    if filter_value and FILTER_COLUMN_NAME in df.columns:
        original_len = len(df)
        df = df[
            df[FILTER_COLUMN_NAME]
            .astype(str)
            .str.strip()
            .str.lower() == filter_value
        ]
        logger.info(f"Filtered '{sheet_name}' for '{filter_value}' → {len(df)} rows (from {original_len})")
    else:
        if filter_value:
            logger.warning(f"Filter column '{FILTER_COLUMN_NAME}' not found in '{sheet_name}'. Skipping filter.")

    return df


def normalize_header(header):
    """
    Normalizes a header string to allow for consistent matching across different sources (e.g., Excel and Word).

    This function:
    - Ensures the header is a string.
    - Applies Unicode normalization to decompose characters (e.g., accented characters).
    - Removes special and control characters such as line breaks, non-breaking spaces, and hidden Unicode marks.
    - Collapses multiple spaces into one.
    - Converts the entire header to lowercase for case-insensitive comparison.

    Args:
        header (str or any): The header to normalize.

    Returns:
        str: A cleaned, lowercase, single-line version of the header suitable for matching.
    """
    # Convert to string if not already
    if not isinstance(header, str):
        header = str(header)

    # Normalize Unicode characters to their canonical decomposition
    header = unicodedata.normalize("NFKD", header)

    # Replace various invisible or problematic Unicode characters with a space or remove them
    header = header.replace('\r', ' ').replace('\n', ' ').replace('\u2028', ' ') \
                   .replace('\u2029', ' ').replace('\u00A0', ' ').replace('\u200B', '') \
                   .replace('\u2060', '').replace('\u200F', '')

    # Remove any remaining control characters (Unicode category starting with 'C')
    header = ''.join(ch for ch in header if unicodedata.category(ch)[0] != 'C')

    # Strip leading/trailing whitespace, collapse multiple spaces into one, convert to lowercase
    return re.sub(r'\s+', ' ', header.strip().lower())


def get_table_after_paragraph(doc, target_para):
    found = False
    target_element = target_para._element
    for block in doc.element.body.iterchildren():
        if block == target_element:
            found = True
            continue
        if found and block.tag.endswith('tbl'):
            return docx.table.Table(block, doc)
    return None


def is_merged_header(table):
    if len(table.rows) < 2:
        return False
    row0, row1 = table.rows[0], table.rows[1]
    if len(row0.cells) != len(row1.cells):
        return True
    merged_texts = ["Observed Results for Process Validation Lots"]
    return any(text in cell.text for cell in row0.cells for text in merged_texts)


# def apply_vac_special_handling(table, df):
#     """
#     Applies VAC special handling on a Word table by copying relevant columns 
#     from the Excel DataFrame to the Word table cells.

#     Rules:
#     - Only the lower header row is used to identify and insert data.
#     - The merged upper cell labeled “Observed Results for Process Validation Lots” is ignored.
#     - Lot numbers (columns) are extracted from Excel between headers:
#         "Validation Acceptance Criteria" and "Meets Validation Acceptance Criteria (Yes/No)"
#     - Only first 3 intermediate columns from Excel are copied, regardless of Word columns.
#     - Copies both headers and data rows from Excel to Word.

#     Args:
#         table: docx.table.Table object representing the Word table.
#         df: pandas.DataFrame containing the Excel data.
#     """

#     from docx.oxml import OxmlElement

#     def clear_cell(cell):
#         """Clears the content of a docx table cell, preserving paragraph formatting."""
#         cell.text = ""

#     logger.info("Starting VAC special handling with hardcoded 3 columns...")

#     # Normalize Word and Excel headers for matching
#     word_headers = [normalize_header(cell.text.strip()) for cell in table.rows[1].cells]
#     df.columns = [normalize_header(col) for col in df.columns]

#     try:
#         vac_word_idx = word_headers.index(VAC_HEADER_START)
#         meets_vac_word_idx = word_headers.index(VAC_HEADER_END)
#         logger.info(f"Found VAC start index in Word: {vac_word_idx}, end index: {meets_vac_word_idx}")
#     except ValueError:
#         logger.warning("VAC headers not found in Word table. Skipping VAC handling.")
#         return

#     try:
#         vac_excel_idx = df.columns.get_loc(VAC_HEADER_START)
#         meets_vac_excel_idx = df.columns.get_loc(VAC_HEADER_END)
#         logger.info(f"Found VAC in Excel columns: from index {vac_excel_idx} to {meets_vac_excel_idx}")
#     except KeyError:
#         logger.warning("VAC headers not found in Excel DataFrame. Skipping VAC handling.")
#         return

#     # Extract intermediate columns from Excel (strictly first 3 only)
#     intermediate_excel_cols = list(df.columns[vac_excel_idx + 1: meets_vac_excel_idx])
#     logger.info(f"All Excel intermediate VAC columns detected: {intermediate_excel_cols}")

#     # Hardcode number of columns to 3 for this logic
#     expected_intermediate_cols = 3
#     if len(intermediate_excel_cols) < expected_intermediate_cols:
#         logger.warning(f"Excel has fewer than {expected_intermediate_cols} intermediate columns. Using available columns.")
#         expected_intermediate_cols = len(intermediate_excel_cols)

#     intermediate_excel_cols = intermediate_excel_cols[:expected_intermediate_cols]
#     logger.info(f"Using first {expected_intermediate_cols} Excel VAC columns: {intermediate_excel_cols}")

#     data_start_row = 2  # Assuming first two rows are headers

#     # Copy Excel headers into Word header row
#     header_row = table.rows[1]
#     for i in range(expected_intermediate_cols):
#         col_idx = vac_word_idx + 1 + i
#         cell = header_row.cells[col_idx]
#         try:
#             clear_cell(cell)
#             header_text = intermediate_excel_cols[i]
#             cell.text = header_text
#             actual_text = cell.text
#             if actual_text != header_text:
#                 logger.warning(f"Mismatch writing header to Word cell[{col_idx}]: expected '{header_text}', got '{actual_text}'")
#             else:
#                 logger.info(f"Copied header '{header_text}' to Word header cell[{col_idx}] successfully")
#         except Exception as e:
#             logger.error(f"Error writing header to Word cell[{col_idx}]: {e}")

#     # Copy Excel data into Word rows
#     max_rows = min(len(df), len(table.rows) - data_start_row)
#     logger.info(f"Copying {max_rows} rows of data from Excel to Word")

#     for row_offset in range(max_rows):
#         word_row = table.rows[data_start_row + row_offset]
#         logger.info(f"Inserting row {row_offset} data from Excel into Word row {data_start_row + row_offset}")

#         for i in range(expected_intermediate_cols):
#             col_idx = vac_word_idx + 1 + i
#             excel_col = intermediate_excel_cols[i]
#             val = str(df.iloc[row_offset][excel_col]) if pd.notna(df.iloc[row_offset][excel_col]) else ""
#             cell = word_row.cells[col_idx]
#             try:
#                 clear_cell(cell)
#                 cell.text = val
#                 actual_text = cell.text
#                 if actual_text != val:
#                     logger.warning(f"Mismatch writing to Word cell[{col_idx}] at row {data_start_row + row_offset}: expected '{val}', got '{actual_text}'")
#                 else:
#                     logger.info(f"Excel[{row_offset}]['{excel_col}'] = '{val}' → Word cell[{col_idx}] successfully")
#             except Exception as e:
#                 logger.error(f"Error writing to Word cell[{col_idx}] at row {data_start_row + row_offset}: {e}")

#     logger.info("VAC special handling completed successfully with hardcoded columns.")


def apply_vac_special_handling(table, df):
    """
    Applies special handling for the VAC section of a Word table by copying intermediate column headers
    and values from an Excel DataFrame, strictly by position (not name matching).
    """
    logger.info("Starting VAC special handling (positional logic only)...")

    # Normalize Word table headers (row 1 assumed as header row)
    word_headers = [normalize_header(cell.text.strip()) for cell in table.rows[1].cells]
    df.columns = [normalize_header(col) for col in df.columns]
    logger.info(f"Normalized {len(word_headers)} Word headers and {len(df.columns)} Excel headers")

    # Step 1: Locate VAC boundaries in Word
    try:
        vac_word_idx = word_headers.index(VAC_HEADER_START)
        meets_vac_word_idx = word_headers.index(VAC_HEADER_END)
        logger.info(f"Word VAC section: Start at {vac_word_idx}, End at {meets_vac_word_idx}")
    except ValueError:
        logger.warning("VAC headers not found in Word table. Skipping VAC special handling.")
        return

    word_slots = list(range(vac_word_idx + 1, meets_vac_word_idx))
    logger.info(f"Word table has {len(word_slots)} slots between VAC columns.")

    # Step 2: Locate VAC boundaries in Excel
    try:
        vac_excel_idx = df.columns.get_loc(VAC_HEADER_START)
        meets_vac_excel_idx = df.columns.get_loc(VAC_HEADER_END)
        logger.info(f"Excel VAC section: Start at {vac_excel_idx}, End at {meets_vac_excel_idx}")
    except KeyError:
        logger.warning("VAC headers not found in Excel DataFrame. Skipping VAC special handling.")
        return

    intermediate_excel_cols = list(df.columns[vac_excel_idx + 1: meets_vac_excel_idx])
    logger.info(f"Excel has {len(intermediate_excel_cols)} intermediate VAC columns: {intermediate_excel_cols}")

    if not intermediate_excel_cols:
        logger.info("No Excel columns found between VAC headers. Skipping.")
        return

    num_to_copy = min(len(intermediate_excel_cols), len(word_slots))
    logger.info(f"Copying {num_to_copy} columns from Excel into Word table.")

    # Step 3: Copy Excel headers into Word table header row
    header_row = table.rows[1]
    for i in range(num_to_copy):
        word_col_idx = word_slots[i]
        header_text = intermediate_excel_cols[i]
        header_row.cells[word_col_idx].text = header_text
        logger.info(f"Inserted header '{header_text}' at Word column index {word_col_idx}")

    # Step 4: Copy Excel data values into Word rows — **always overwrite**
    data_start_row = 2
    max_rows = min(len(df), len(table.rows) - data_start_row)
    logger.info(f"Preparing to copy data for {max_rows} rows...")

    for row_offset in range(max_rows):
        word_row = table.rows[data_start_row + row_offset]
        for i in range(num_to_copy):
            word_col_idx = word_slots[i]
            excel_col = intermediate_excel_cols[i]
            value = str(df.iloc[row_offset][excel_col]) if pd.notna(df.iloc[row_offset][excel_col]) else ""
            cell = word_row.cells[word_col_idx]
            cell.text = value
            logger.info(f"Row {data_start_row + row_offset}, Col {word_col_idx}: Inserted '{value}'")

    logger.info("VAC special handling completed successfully (positional logic).\n")


def build_column_mapping(word_headers, excel_headers, partial_match_rules: Optional[list[str]] = None):
    """
    Builds a mapping between Word table column indexes and Excel DataFrame column indexes
    based on exact and partial header matches.

    Args:
        word_headers (list[str]): Normalized headers from Word table.
        excel_headers (list[str]): Normalized headers from Excel DataFrame.
        partial_match_rules (list[str], optional): List of rules for fuzzy/partial matching.

    Returns:
        dict: Mapping from Word column index to Excel column index.
    """
    mapping = {}

    # Normalize headers
    normalized_excel_headers = [normalize_header(h) for h in excel_headers]
    normalized_word_headers = [normalize_header(h) for h in word_headers]

    for w_idx, word_header in enumerate(normalized_word_headers):
        matched = False

        # Try exact match first
        if word_header in normalized_excel_headers:
            mapping[w_idx] = normalized_excel_headers.index(word_header)
            continue

        # Try partial/fuzzy matching using rules (if provided)
        for e_idx, excel_header in enumerate(normalized_excel_headers):
            for rule in partial_match_rules or []:
                rule_norm = normalize_header(rule)
                if rule_norm in word_header and rule_norm in excel_header:
                    mapping[w_idx] = e_idx
                    matched = True
                    break
            if matched:
                break

    return mapping


def insert_table_after_paragraph(doc, placeholder: str, df) -> bool:
    """
    Finds a paragraph with the given placeholder text and fills the table
    immediately following it using the provided DataFrame.

    Returns True if successful, False otherwise.
    """
    norm_marker = normalize_header(placeholder)

    for idx, para in enumerate(doc.paragraphs):
        para_text = normalize_header(para.text)

        if norm_marker in para_text:
            logger.info(f"Found placeholder '{placeholder}' in paragraph {idx}")

            # Optional: Remove or replace the placeholder text
            para.text = para.text.replace(placeholder, "")  # or replace with a proper title

            table = get_table_after_paragraph(doc, para)
            if table:
                fill_table_by_headers(table, df, partial_match_rules=PARTIAL_MATCH_HEADERS)
                logger.info(f"Inserted table after placeholder: {placeholder}")
                return True
            else:
                logger.warning(f"No table found after placeholder: {placeholder}")
                return False

    logger.warning(f"Placeholder paragraph '{placeholder}' not found in document.")
    return False


def fill_table_by_headers(table, df, partial_match_rules: Optional[list[str]] = None):
    if len(table.rows) < 2:
        logger.warning("Table has fewer than 2 rows. Skipping.")
        return

    # Keep original headers before normalization for case preservation
    original_excel_cols = df.columns.tolist()
    df.columns = [normalize_header(col) for col in original_excel_cols]
    apply_vac_special_handling(table, df)

    # Check for merged header spanning 2 rows
    has_merged_header = is_merged_header(table)
    row0 = table.rows[0].cells
    row1 = table.rows[1].cells if len(table.rows) > 1 else []

    combined_headers = []
    for col_idx in range(len(row0)):
        top = row0[col_idx].text.strip() if col_idx < len(row0) else ""
        bottom = row1[col_idx].text.strip() if col_idx < len(row1) else ""
        combined = bottom if has_merged_header else (bottom or top)
        combined_headers.append(combined)

    # Normalize combined headers from Word table for mapping
    word_headers = [normalize_header(h) for h in combined_headers]

    # Build mapping from Word header idx to Excel df column idx
    col_mapping = build_column_mapping(word_headers, df.columns, partial_match_rules)
    if not col_mapping:
        logger.warning("No matching headers found between Excel and Word table.")
        return

    # Write original headers (preserving case) into Word table header row(s)
    for w_col_idx in col_mapping:
        df_col_idx = col_mapping[w_col_idx]
        header_text = original_excel_cols[df_col_idx]
        # If header spans two rows, write to correct row
        target_row = 1 if has_merged_header else 0
        table.rows[target_row].cells[w_col_idx].text = header_text

    # Set start row for data fill based on merged header presence
    DATA_START_INDEX = 2 if has_merged_header else 1
    rows_to_fill = min(len(df), len(table.rows) - DATA_START_INDEX)

    # Fill data cells from Excel into Word table
    for i in range(rows_to_fill):
        for w_col_idx, df_col_idx in col_mapping.items():
            cell = table.rows[i + DATA_START_INDEX].cells[w_col_idx]
            if not cell.text.strip():
                # cell.text = str(df.iat[i, df_col_idx])
                value = df.iat[i, df_col_idx]
                if pd.isna(value):
                    value = ""
                cell.text = str(value)

    logger.info("Table populated successfully.")

# ------------------------------------------------------------------------------
# ✅ MAIN ENTRY FUNCTION
# ------------------------------------------------------------------------------

def populate_tables(bucket_name: str,
                    rules_s3_key: str,
                    source_xls_s3_key: str,
                    target_docx_s3_key: str,
                    tables_mapping: dict):
    """
    Populate tables into DOCX file stored on S3 using mapping rules and source Excel tables also in S3.

    Args:
        bucket_name (str): S3 bucket name.
        rules_s3_key (str): S3 key of the Excel file that contains mapping rules (sheet names, placeholders).
        source_xls_s3_key (str): S3 key of the Excel file containing actual source data tables.
        target_docx_s3_key (str): S3 key of the target DOCX report file to update.
        tables_mapping (dict): Dictionary mapping placeholders to sheet names or mapping info.
    """
    try:
        logger.info("Starting populate_tables ----------------->")
        # 1. Load rules Excel from S3 (currently unused but kept for future)
        logger.info(f"Loading mapping rules Excel from S3: {rules_s3_key}")
        _ = s3.get_object(Bucket=bucket_name, Key=rules_s3_key)  # Not used directly here

        # 2. Load target DOCX from S3
        logger.info(f"Loading target DOCX report from S3: {target_docx_s3_key}")
        response = s3.get_object(Bucket=bucket_name, Key=target_docx_s3_key)
        docx_bytes = response['Body'].read()
        doc = Document_docx(io.BytesIO(docx_bytes))

        # 3. Loop through each placeholder and insert corresponding table
        for placeholder, sheet_name in tables_mapping.items():
            logger.info(f"Processing placeholder '{placeholder}' with sheet '{sheet_name}'")

            # Load and filter data
            df = load_and_conditionally_filter_excel_sheet(
                bucket_name=bucket_name,
                xls_s3_key=source_xls_s3_key,
                sheet_name=sheet_name,
                placeholder=placeholder
            )

            # Insert table in the DOCX
            inserted = insert_table_after_paragraph(doc, placeholder, df)
            if not inserted:
                logger.warning(f"Placeholder '{placeholder}' not found in document.")

        # 4. Save the updated DOCX back to S3
        output_stream = io.BytesIO()
        doc.save(output_stream)
        output_stream.seek(0)
        s3.put_object(Bucket=bucket_name, Key=target_docx_s3_key, Body=output_stream.getvalue())
        logger.info(f"Updated DOCX uploaded back to S3 at {target_docx_s3_key}")

    except Exception as e:
        logger.error(f"Error in populate_tables: {e}")
        raise